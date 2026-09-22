#!/usr/bin/env python3
"""Build the A4 poster as an editable PowerPoint (.pptx).

Content is read from `mouse2human-poster.html` (single source of truth);
the PPTX mirrors its header, key-number band, two text columns, figure
strip and footer with real, editable text boxes.

    python build_poster_pptx.py
"""
from __future__ import annotations

import math
import pathlib
import re

from bs4 import BeautifulSoup
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = pathlib.Path(__file__).parent
HTML = HERE / "mouse2human-poster.html"
OUT = HERE / "mouse2human-poster.pptx"

INK, BODY, MUTED, SOFT = "000000", "1B2130", "5B6377", "9AA3B8"
CYAN, BLUE, PERI, RULE = "4DBBD5", "4361EE", "A1B0F7", "D9DEEC"
DARK, CARD2 = "0B0E16", "F6F8FD"
SERIF, SANS = "Merriweather", "Arial"

SW, SH = 8.2677, 11.6929        # A4 portrait, inches
MX = 0.33
CW = SW - 2 * MX
COLW = (CW - 0.30) / 2
RX = MX + COLW + 0.30


def box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return tb, tf


def para(tf, first=False, spacing=1.2, after=0, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.line_spacing = spacing
    p.space_after = Pt(after)
    p.alignment = align
    return p


def put(p, text, size=6, font=SANS, color=BODY, bold=False, italic=False, spc=None):
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    if spc is not None:
        r.font._rPr.set("spc", str(int(spc * 100)))
    return r


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.5):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = RGBColor.from_string(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def est_lines(text, w, size, factor=0.52):
    cpl = max(6, int(w * 72 / (size * factor)))
    return max(1, math.ceil(len(text) / cpl))


def rich_runs(el):
    out = []

    def walk(node, bold=False, italic=False):
        for ch in node.children:
            if isinstance(ch, str):
                t = re.sub(r"\s+", " ", ch)
                if t:
                    out.append((t, bold, italic))
            elif ch.name in ("strong", "b"):
                walk(ch, True, italic)
            elif ch.name in ("em", "i"):
                walk(ch, bold, True)
            elif ch.name in ("code", "span"):
                walk(ch, bold, italic)
            elif ch.name == "br":
                out.append(("\n", bold, italic))
            else:
                walk(ch, bold, italic)

    walk(el)
    while out and not out[0][0].strip():
        out.pop(0)
    while out and not out[-1][0].strip():
        out.pop()
    return out


def picture_fit(slide, path, x, y, max_w, max_h, center=True):
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = iw / ih
    w = max_w
    h = w / ratio
    if h > max_h:
        h = max_h
        w = h * ratio
    left = x + (max_w - w) / 2 if center else x
    slide.shapes.add_picture(str(path), Inches(left), Inches(y), Inches(w), Inches(h))
    return y + h


def main() -> None:
    soup = BeautifulSoup(HTML.read_text(encoding="utf-8"), "lxml")
    header = soup.find("header")

    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ── header ──────────────────────────────────────────────────────────────
    H = 1.62
    rect(slide, 0, 0, SW, H, fill=DARK)
    rect(slide, 0, 0, SW, 0.055, fill=CYAN)
    rect(slide, SW / 3, 0, SW / 3, 0.055, fill=BLUE)
    rect(slide, 2 * SW / 3, 0, SW / 3, 0.055, fill=PERI)

    tb, tf = box(slide, MX, 0.16, CW, 0.2)
    p = para(tf, True, spacing=1.0)
    put(p, header.select_one(".eyebrow").get_text(" ", strip=True).upper(),
        size=5.6, color="9AA3B8", spc=1.6)

    tb, tf = box(slide, MX, 0.32, CW, 0.5)
    put(para(tf, True), header.find("h1").get_text(" ", strip=True),
        size=22, font=SERIF, color="FFFFFF", bold=True)

    tb, tf = box(slide, MX, 0.78, CW, 0.4)
    p = para(tf, True, spacing=1.15)
    put(p, header.select_one(".sub").get_text(" ", strip=True), size=9, font=SERIF, color="D8DEEC")

    tb, tf = box(slide, MX, 1.10, CW, 0.3)
    p = para(tf, True, spacing=1.3)
    put(p, header.select_one(".authors").get_text(" ", strip=True), size=5.6, color="E6E9F2")

    tb, tf = box(slide, MX, 1.34, CW, 0.3)
    p = para(tf, True, spacing=1.3)
    put(p, header.select_one(".affil").get_text(" ", strip=True), size=5.4, color="9AA3B8")
    tb, tf = box(slide, MX, 1.47, CW, 0.2)
    p = para(tf, True, spacing=1.2)
    put(p, header.select_one(".repo").get_text(" ", strip=True), size=5.2, color=PERI)

    # ── key-number band ─────────────────────────────────────────────────────
    y = H
    band = soup.select(".band > div")
    bw = CW / len(band)
    rect(slide, MX, y, CW, 0.48, fill="FFFFFF", line=RULE)
    for i, d in enumerate(band):
        x = MX + i * bw
        if i:
            rect(slide, x, y + 0.06, 0.006, 0.36, fill=RULE)
        tb, tf = box(slide, x + 0.06, y + 0.05, bw - 0.12, 0.2)
        p = para(tf, True, spacing=1.0)
        put(p, d.select_one(".n").get_text(" ", strip=True), size=13, font=SERIF,
            color=INK if "ink" in (d.select_one(".n").get("class") or []) else BLUE)
        tb, tf = box(slide, x + 0.06, y + 0.25, bw - 0.12, 0.2)
        p = para(tf, True, spacing=1.15)
        put(p, d.select_one(".l").get_text(" ", strip=True).upper(), size=4.8,
            color=MUTED, spc=0.9)
    y += 0.56

    # ── two text columns ────────────────────────────────────────────────────
    BOTTOM = 9.86
    strip_h = 1.30
    avail_bottom = SH - strip_h - 0.34          # footer + strip space
    col_bottom = avail_bottom

    ui = [s for s in soup.select("main section")]
    items = []
    for sec in ui:
        title = sec.find("h2").get_text(" ", strip=True)
        blocks = [("h2", title)]
        for el in sec.find_all(recursive=False):
            if el.name == "h2":
                continue
            cls = el.get("class") or []
            if el.name == "p":
                blocks.append(("p", el))
            elif el.name == "ul":
                blocks.append(("ul", el))
            elif "panel" in cls:
                blocks.append(("panel", el))
            elif "cards" in cls:
                blocks.append(("cards", el))
        items.append(blocks)

    def h_block(kind, payload, w):
        if kind == "h2":
            return 0.16 + est_lines(payload, w, 8.2) * 8.2 * 1.1 / 72
        if kind == "p":
            return est_lines(payload.get_text(" ", strip=True), w, 6.1) * 6.1 * 1.36 / 72 + 0.05
        if kind == "ul":
            n = len(payload.find_all("li", recursive=False))
            lines = sum(est_lines(li.get_text(" ", strip=True), w - 0.12, 6.1) for li in payload.find_all("li", recursive=False))
            return lines * 6.1 * 1.32 / 72 + n * 0.018 + 0.05
        if kind == "panel":
            big = payload.select_one(".big")
            pt = payload.select_one(".pt")
            ps = payload.find_all("p")
            hh = 0.05 + (0.10 if pt else 0) + (0.18 if big else 0)
            for p in ps:
                if p is pt or p is big:
                    continue
                hh += est_lines(p.get_text(" ", strip=True), w, 6.1) * 6.1 * 1.36 / 72 + 0.04
            return hh + 0.06
        if kind == "cards":
            hh = 0.0
            for c in payload.find_all("div", class_="card", recursive=False):
                hh = max(hh, 0.62)
            return hh + 0.06
        return 0.1

    col = 0
    cx = [MX, RX]
    cy = y + 0.08
    for blocks in items:
        total = sum(h_block(k, v, COLW) for k, v in blocks) + 0.20
        if cy + total > col_bottom and col == 0:
            col, cy = 1, y + 0.08
        x = cx[col]
        for kind, payload in blocks:
            if kind == "h2":
                tb, tf = box(slide, x, cy, COLW, 0.2)
                p = para(tf, True, spacing=1.0)
                m = re.match(r"^(\d+)", payload)
                if m:
                    put(p, payload[:len(m.group(1))], size=8.2, font=SERIF, color=CYAN, bold=True)
                    put(p, payload[len(m.group(1)):], size=8.2, font=SERIF, color=INK, bold=True)
                else:
                    put(p, payload, size=8.2, font=SERIF, color=INK, bold=True)
                cy += 0.13
                rect(slide, x, cy, COLW, 0.018, fill=INK)
                cy += 0.05
            elif kind == "p":
                tb, tf = box(slide, x, cy, COLW, 0.2)
                p = para(tf, True, spacing=1.34)
                for t, b, i in rich_runs(payload):
                    put(p, t, size=6.1, color=INK if b else BODY, bold=b, italic=i)
                cy += h_block(kind, payload, COLW)
            elif kind == "ul":
                tb, tf = box(slide, x, cy, COLW, 0.2)
                first = True
                for li in payload.find_all("li", recursive=False):
                    p = para(tf, first=first, spacing=1.3, after=1.2)
                    first = False
                    put(p, "•  ", size=6.1, color=PERI, bold=True)
                    for t, b, i in rich_runs(li):
                        put(p, t, size=6.1, color=INK if b else BODY, bold=b, italic=i)
                cy += h_block(kind, payload, COLW)
            elif kind == "panel":
                strong = "strong" in (payload.get("class") or [])
                rect(slide, x, cy, COLW, 0.018, fill=BLUE if strong else PERI)
                cy += 0.06
                pt = payload.select_one(".pt")
                big = payload.select_one(".big")
                if pt:
                    tb, tf = box(slide, x, cy, COLW, 0.15)
                    p = para(tf, True, spacing=1.1)
                    put(p, pt.get_text(" ", strip=True).upper(), size=4.8, color=MUTED, spc=0.9)
                    cy += 0.11
                if big:
                    tb, tf = box(slide, x, cy, COLW, 0.2)
                    p = para(tf, True, spacing=1.0)
                    put(p, big.get_text(" ", strip=True), size=10.5, font=SERIF,
                        color=BLUE if strong else "2E9DB7")
                    cy += 0.18
                for p_el in payload.find_all("p"):
                    if p_el is pt or p_el is big:
                        continue
                    tb, tf = box(slide, x, cy, COLW, 0.2)
                    p = para(tf, True, spacing=1.34)
                    for t, b, i in rich_runs(p_el):
                        put(p, t, size=6.1, color=INK if b else BODY, bold=b, italic=i)
                    cy += est_lines(p_el.get_text(" ", strip=True), COLW, 6.1) * 6.1 * 1.34 / 72 + 0.045
                cy += 0.05
            elif kind == "cards":
                cards = payload.find_all("div", class_="card", recursive=False)
                gap = 0.06
                cwid = (COLW - gap * 2) / 3
                for j, c in enumerate(cards):
                    cxx = x + j * (cwid + gap)
                    rect(slide, cxx, cy, cwid, 0.62, fill="FFFFFF", line=RULE)
                    tb, tf = box(slide, cxx + 0.05, cy + 0.05, cwid - 0.1, 0.12)
                    p = para(tf, True, spacing=1.0)
                    put(p, c.select_one(".tag").get_text(" ", strip=True).upper(),
                        size=4.2, color=CYAN, spc=0.9)
                    tb, tf = box(slide, cxx + 0.05, cy + 0.16, cwid - 0.1, 0.16)
                    p = para(tf, True, spacing=1.05)
                    put(p, c.find("h3").get_text(" ", strip=True), size=5.6, font=SERIF, color=INK)
                    tb, tf = box(slide, cxx + 0.05, cy + 0.28, cwid - 0.1, 0.26)
                    p = para(tf, True, spacing=1.22)
                    put(p, c.find("p").get_text(" ", strip=True), size=4.5, color=MUTED)
                    tb, tf = box(slide, cxx + 0.05, cy + 0.53, cwid - 0.1, 0.1)
                    p = para(tf, True, spacing=1.1)
                    put(p, c.select_one(".url").get_text(" ", strip=True), size=4.4, color=BLUE)
                cy += 0.68
            cy += 0.105

    # ── figure strip ────────────────────────────────────────────────────────
    sy = avail_bottom + 0.06
    rect(slide, 0, sy - 0.04, SW, SH - sy + 0.04, fill=CARD2)
    rect(slide, 0, sy - 0.04, SW, 0.008, fill=RULE)
    figs = soup.select(".strip figure")
    fw = (CW - 3 * 0.12) / 4
    for i, f in enumerate(figs):
        fx = MX + i * (fw + 0.12)
        src = HERE / f.find("img")["src"]
        picture_fit(slide, src, fx, sy + 0.04, fw, 0.72)
        tb, tf = box(slide, fx, sy + 0.80, fw, 0.3)
        p = para(tf, True, spacing=1.2)
        cap = f.find("figcaption")
        parts = rich_runs(cap)
        for t, b, it in parts:
            put(p, t, size=4.3, color=INK if b else MUTED, bold=b, italic=it)

    # ── footer ──────────────────────────────────────────────────────────────
    fy = SH - 0.20
    rect(slide, 0, fy - 0.04, SW, 0.008, fill=RULE)
    tb, tf = box(slide, MX, fy, CW * 0.5, 0.15)
    p = para(tf, True, spacing=1.0)
    put(p, "Mouse2Human · multilayer — poster (A4) · September 2026", size=5.0, color=MUTED)
    tb, tf = box(slide, MX + CW * 0.5, fy, CW * 0.5, 0.15)
    p = para(tf, True, spacing=1.0, align=PP_ALIGN.RIGHT)
    put(p, "Figures from the manuscript · Genes and Immunity (under review)", size=5.0, color=BLUE)

    prs.core_properties.title = "From mice to humans — poster (A4)"
    prs.core_properties.author = "Wasim Aluísio Prates-Syed · Cabral-Miranda Lab"
    prs.save(OUT)
    print(f"wrote {OUT.name}")


if __name__ == "__main__":
    main()
